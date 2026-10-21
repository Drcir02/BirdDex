#!/usr/bin/env python3
"""Adds an `info` object to every bird in public/data/species.json: when, where and how to find it.

Run it from the project root, after scripts/build-species.py (that script writes species.json without `info`):

    .venv/bin/python scripts/build-species-info.py --avonet AVONET.xlsx --elton BirdFuncDat.txt [--review FILE] [--no-write]

Inputs
  * public/data/species.json and public/models/birdnet-v2.4-geo-fp16.tflite (the same location model build-species.py uses)
  * AVONET Supplementary dataset 1.xlsx (Tobias et al. 2022, Ecology Letters, doi:10.1111/ele.13898, CC BY 4.0), sheets
    AVONET1_BirdLife and "BirdLife-eBird crosswalk". Download: https://figshare.com/articles/dataset/16586228
  * BirdFuncDat.txt (EltonTraits 1.0, Wilman et al. 2014, Ecology 95:2027, doi:10.1890/13-1917.1).
    Download: https://esapubs.org/archive/ecol/E095/178/BirdFuncDat.txt (MD5 d6197b2cd90ca3ece0a7393abbf8b7fc)

What each field is made from
  months     BirdNET's location model, which learned from bird observation records. For a bird of Czechia: the average over the
             Czech sample points of each week's score. For a bird that is only in Europe: the average over the European
             points where the model puts the bird at all (the range of the bird). Four weeks make one month, and each month is
             compared to the best month of the same bird: 3 = most likely, 2 = likely, 1 = a few, 0 = not there.
  countries  Birds that are not in Czechia only: the countries of the European sample points (POINT_COUNTRY) where the model
             gives the bird a score of 0.03 or more and a quarter of its best country's score, best first (at most 10).
  place      AVONET's habitat class of the bird, written as one short place (PLACE_OF_HABITAT).
  day        When it is most active. EltonTraits says which birds are nocturnal; the rest is an ESTIMATE by family (FAMILY_RULES).
  shy        How close a person gets: 1 lets you come close, 2 is wary, 3 is shy or far away and needs a zoom.
             An ESTIMATE by family and genus (FAMILY_RULES, SPECIES_RULES), not from any dataset.
  tip        Which group text of the app (dex.info.tips.*) belongs to the bird. By family, a few by genus (GENUS_TIP).

shy, day and tip are written by hand, from general knowledge of the families. They should be read through by somebody who
knows birds: --review writes one line per bird for that.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parent.parent
SPECIES_FILE = ROOT / "public" / "data" / "species.json"

_spec = importlib.util.spec_from_file_location("build_species", ROOT / "scripts" / "build-species.py")
bs = importlib.util.module_from_spec(_spec)
sys.modules["build_species"] = bs
_spec.loader.exec_module(bs)  # type: ignore[union-attr]

# Months get 4 of BirdNET's 48 weeks each.
WEEKS_PER_MONTH = 4
# A month counts as level 3, 2 or 1 when its score is at least this share of the bird's best month.
LEVEL_SHARES = ((0.6, 3), (0.2, 2), (0.05, 1))
# The score where BirdNET's own species list counts a bird as present at a place.
PRESENT_SCORE = 0.03
# A country counts when its best point reaches PRESENT_SCORE and this share of the bird's best country: the model spreads
# a little of the score of a seabird or a mountain bird over places where it hardly ever is.
COUNTRY_SHARE = 0.25
# A bird that has no point at all at that score uses its best points instead.
RANGE_MIN_POINTS = 3
MAX_COUNTRIES = 10
# For the few birds that are never above PRESENT_SCORE anywhere, countries still count from this score on.
COUNTRY_FLOOR = 0.005

# ---------------------------------------------------------------------------------------------------------------------
# Countries of the European sample points (the names are those in build-species.py's EUROPE_POINTS).
# ---------------------------------------------------------------------------------------------------------------------

POINT_COUNTRY: dict[str, str] = {}
for _country, _names in {
    "GB": "London|Norfolk coast|Cornwall|Peak District|Cairngorms|Shetland|Orkney|Northumberland coast|Snowdonia",
    "IE": "Dublin|Kerry",
    "IS": "Reykjavik|Myvatn|Vik",
    "NO": "Oslo|Bergen|Trondheim|Tromso|Varanger|Hardangervidda|Dovrefjell",
    "SE": "Stockholm|Gothenburg|Oland|Falsterbo|Abisko|Umea",
    "DK": "Copenhagen|Skallingen",
    "FI": "Helsinki|Oulu|Inari|Kuusamo",
    "EE": "Tallinn|Matsalu",
    "LV": "Riga",
    "LT": "Vilnius|Nemunas delta",
    "RU": "St Petersburg",
    "BY": "Minsk",
    "UA": "Kyiv|Lviv|Odesa|Crimea",
    "MD": "Chisinau",
    "PL": "Bialowieza|Warsaw|Hel|Krakow|Zakopane|Biebrza|Wroclaw",
    "DE": "Berlin|Hamburg|Munich|Black Forest|Frankfurt|Dresden|Ruegen|Bavarian Forest|Cologne|Wadden Sea",
    "NL": "Amsterdam|Texel",
    "BE": "Brussels",
    "LU": "Luxembourg",
    "FR": "Paris|Brittany|Camargue|French Pyrenees|Alsace|Loire|Landes|Corsica|Vanoise|Baie de Somme|Massif Central|Crau",
    "CH": "Zurich|Engadin|Valais",
    "AT": "Vienna|Hohe Tauern|Neusiedler See|Salzkammergut",
    "SK": "Bratislava|High Tatras|Kosice",
    "HU": "Budapest|Hortobagy|Balaton",
    "RO": "Bucharest|Brasov|Cluj|Dobruja",
    "SI": "Ljubljana",
    "HR": "Zagreb|Split|Plitvice|Istria",
    "RS": "Belgrade",
    "BA": "Sarajevo",
    "ME": "Skadar lake",
    "AL": "Tirana",
    "MK": "Skopje",
    "BG": "Sofia|Rila|Burgas",
    "GR": "Athens|Crete|Lesbos|Thessaloniki|Evros|Peloponnese|Corfu|Rhodes|Olympus",
    "TR": "Istanbul|Izmir|Dardanelles|Bodrum",
    "IT": "Rome|Po delta|Milan|Sicily|Sardinia|Gargano|Gran Paradiso|Calabria|Tuscany|Dolomites|Abruzzo|Venice lagoon",
    "ES": "Madrid|Donana|Monfrague|Ordesa|Ebro delta|Picos de Europa|Galicia|Tarifa|Almeria|Mallorca|Bardenas|Gredos|Albufera|Barcelona|Sierra Nevada|Trujillo|Burgos|Belchite",
    "PT": "Ria Formosa|Lisbon|Porto|Castro Verde",
}.items():
    for _name in _names.split("|"):
        assert _name not in POINT_COUNTRY, _name
        POINT_COUNTRY[_name] = _country

# ---------------------------------------------------------------------------------------------------------------------
# place: one short place from AVONET's habitat class
# ---------------------------------------------------------------------------------------------------------------------

PLACE_OF_HABITAT = {
    "Wetland": "water",
    "Riverine": "river",
    "Coastal": "coast",
    "Marine": "sea",
    "Forest": "forest",
    "Woodland": "woods",
    "Shrubland": "scrub",
    "Grassland": "fields",
    "Human Modified": "town",
    "Rock": "rocks",
    "Desert": "steppe",
}

# Where AVONET's class does not fit what the bird is to a visitor in Czechia or Europe: the reason is on each line.
PLACE_OVERRIDES: dict[str, str] = {
    "Pavo cristatus": "woods",  # a park and castle-garden bird
    "Numida meleagris": "fields",  # a farmyard bird
    "Cairina moschata": "water",  # a farm and park duck
    "Chrysolophus pictus": "woods",
    "Chrysolophus amherstiae": "woods",
    "Cygnus atratus": "water",
    "Pandion haliaetus": "water",  # AVONET says marine; in Czechia it fishes on ponds and rivers
    "Haliaeetus albicilla": "water",  # AVONET says rock; the eagles of Czechia fish on the big ponds
    "Aquila chrysaetos": "rocks",
    "Falco peregrinus": "rocks",
    "Falco tinnunculus": "fields",
    "Milvus milvus": "fields",
    "Luscinia svecica": "scrub",
    "Prunella collaris": "rocks",
    "Phoenicurus ochruros": "town",
    "Aix galericulata": "water",
    "Coturnix coturnix": "fields",
    "Phasianus colchicus": "fields",
    "Bubulcus ibis": "fields",
    "Gyps fulvus": "rocks",
    "Halcyon smyrnensis": "water",
    "Hippolais olivetorum": "woods",
    "Psittacula krameri": "town",
}
# A bird of the Czech list that AVONET puts on the coast or at sea is met on ponds and rivers when it comes to Czechia.
INLAND_PLACES = {"coast": "water", "sea": "water"}

# ---------------------------------------------------------------------------------------------------------------------
# shy, day and tip: estimates by family, with exceptions by genus and by species
#   shy 1 = lets you come close, 2 = wary, 3 = shy or far away (needs a zoom)
#   day   = day (all day), dawn (best early in the morning), dawn_dusk (morning and evening), dusk (twilight and night), night
# ---------------------------------------------------------------------------------------------------------------------

# family: (shy, day, tip)
FAMILY_RULES: dict[str, tuple[int, str, str]] = {
    "Anatidae": (2, "day", "waterfowl"),
    "Numididae": (1, "day", "parkbirds"),
    "Phasianidae": (2, "dawn_dusk", "gamebirds"),
    "Columbidae": (1, "day", "pigeons"),
    "Pteroclidae": (3, "day", "gamebirds"),
    "Otididae": (3, "day", "gamebirds"),
    "Cuculidae": (3, "dawn", "cuckoos"),
    "Caprimulgidae": (3, "dusk", "nightjars"),
    "Apodidae": (3, "day", "swifts"),
    "Rallidae": (3, "dawn_dusk", "rails"),
    "Gruidae": (3, "dawn_dusk", "cranes"),
    "Burhinidae": (3, "dusk", "waders"),
    "Recurvirostridae": (2, "day", "waders"),
    "Haematopodidae": (2, "day", "waders"),
    "Charadriidae": (2, "day", "waders"),
    "Scolopacidae": (3, "day", "waders"),
    "Glareolidae": (3, "day", "waders"),
    "Stercorariidae": (3, "day", "seabirds"),
    "Alcidae": (3, "day", "seabirds"),
    "Laridae": (2, "day", "gulls"),
    "Phoenicopteridae": (3, "day", "waterbirds"),
    "Podicipedidae": (2, "day", "grebes"),
    "Gaviidae": (3, "day", "divers"),
    "Hydrobatidae": (3, "day", "seabirds"),
    "Procellariidae": (3, "day", "seabirds"),
    "Ciconiidae": (2, "day", "storks"),
    "Sulidae": (3, "day", "seabirds"),
    "Phalacrocoracidae": (2, "day", "waterbirds"),
    "Threskiornithidae": (3, "day", "waterbirds"),
    "Ardeidae": (3, "dawn_dusk", "herons"),
    "Pandionidae": (3, "day", "raptors"),
    "Accipitridae": (3, "day", "raptors"),
    "Tytonidae": (3, "night", "owls"),
    "Strigidae": (3, "dusk", "owls"),
    "Upupidae": (2, "day", "hoopoe"),
    "Meropidae": (2, "day", "beeeater"),
    "Alcedinidae": (3, "day", "kingfisher"),
    "Halcyonidae": (3, "day", "kingfisher"),
    "Coraciidae": (3, "day", "roller"),
    "Picidae": (2, "dawn", "woodpeckers"),
    "Falconidae": (3, "day", "raptors"),
    "Psittaculidae": (2, "day", "escapees"),
    "Psittacidae": (2, "day", "escapees"),
    "Oriolidae": (3, "dawn", "oriole"),
    "Laniidae": (2, "day", "shrikes"),
    "Corvidae": (2, "day", "corvids"),
    "Paridae": (2, "day", "smallwoods"),
    "Remizidae": (3, "day", "reedbed"),
    "Alaudidae": (2, "dawn", "openground"),
    "Panuridae": (3, "dawn", "reedbed"),
    "Cisticolidae": (3, "day", "reedbed"),
    "Acrocephalidae": (3, "dawn", "reedbed"),
    "Locustellidae": (3, "dusk", "reedbed"),
    "Hirundinidae": (2, "day", "swallows"),
    "Phylloscopidae": (2, "dawn", "warblers"),
    "Cettiidae": (3, "dawn", "reedbed"),
    "Aegithalidae": (2, "day", "smallwoods"),
    "Sylviidae": (2, "dawn", "warblers"),
    "Leiothrichidae": (2, "day", "escapees"),
    "Regulidae": (2, "day", "smallwoods"),
    "Tichodromidae": (3, "day", "wallcreeper"),
    "Sittidae": (2, "day", "smallwoods"),
    "Certhiidae": (2, "day", "smallwoods"),
    "Troglodytidae": (2, "dawn", "smallwoods"),
    "Cinclidae": (2, "day", "dipper"),
    "Sturnidae": (1, "day", "starlings"),
    "Turdidae": (2, "dawn", "thrushes"),
    "Muscicapidae": (2, "dawn", "thrushes"),
    "Bombycillidae": (2, "day", "waxwing"),
    "Estrildidae": (2, "day", "escapees"),
    "Prunellidae": (2, "dawn", "smallwoods"),
    "Passeridae": (1, "day", "sparrows"),
    "Motacillidae": (2, "day", "openground"),
    "Fringillidae": (2, "day", "finches"),
    "Calcariidae": (2, "day", "openground"),
    "Emberizidae": (2, "dawn", "finches"),
}

# genus: (shy or None, day or None, tip or None): replaces the parts that are not None
GENUS_RULES: dict[str, tuple[int | None, str | None, str | None]] = {
    "Anser": (3, "dawn_dusk", "geese"),
    "Branta": (3, "dawn_dusk", "geese"),
    "Cygnus": (2, None, None),
    "Mergus": (3, None, None),
    "Mergellus": (3, None, None),
    "Bucephala": (3, None, None),
    "Somateria": (3, None, None),
    "Melanitta": (3, None, None),
    "Clangula": (3, None, None),
    "Histrionicus": (3, None, None),
    "Tetrao": (3, "dawn", "grouse"),
    "Lyrurus": (3, "dawn", "grouse"),
    "Tetrastes": (3, "dawn", "grouse"),
    "Lagopus": (3, "day", "grouse"),
    "Fulica": (1, "day", "coots"),
    "Gallinula": (1, "day", "coots"),
    "Porphyrio": (1, "day", "coots"),
    "Crex": (3, "dusk", None),
    "Rallus": (3, "dusk", None),
    "Porzana": (3, "dusk", None),
    "Zapornia": (3, "dusk", None),
    "Scolopax": (3, "dusk", None),
    "Gallinago": (3, "dawn_dusk", None),
    "Lymnocryptes": (3, "dawn_dusk", None),
    "Accipiter": (3, None, None),
    "Chrysolophus": (2, "day", "parkbirds"),
    "Pavo": (1, "day", "parkbirds"),
    "Delichon": (1, None, None),
    "Hirundo": (1, None, None),
    "Ficedula": (2, None, None),
    "Oenanthe": (2, "day", None),
    "Saxicola": (2, "day", None),
    "Luscinia": (3, "dusk", None),
    "Locustella": (3, "dusk", None),
    "Pyrrhocorax": (1, None, None),
    "Montifringilla": (1, None, None),
}

# species: (shy or None, day or None)
SPECIES_RULES: dict[str, tuple[int | None, str | None]] = {
    "Anas platyrhynchos": (1, None),
    "Cygnus olor": (1, None),
    "Cygnus atratus": (1, None),
    "Branta canadensis": (1, None),
    "Anser anser": (2, None),
    "Cairina moschata": (1, None),
    "Aix galericulata": (2, None),
    "Columba oenas": (2, None),
    "Ciconia ciconia": (1, None),
    "Ciconia nigra": (3, None),
    "Ardea cinerea": (2, None),
    "Buteo buteo": (2, None),
    "Falco tinnunculus": (2, None),
    "Chroicocephalus ridibundus": (1, None),
    "Larus michahellis": (1, None),
    "Caprimulgus europaeus": (None, "dusk"),  # EltonTraits calls nightjars nocturnal; they are seen in the twilight
    "Caprimulgus ruficollis": (None, "dusk"),
    "Athene noctua": (2, "dusk"),
    "Glaucidium passerinum": (3, "dawn_dusk"),
    "Surnia ulula": (3, "day"),
    "Bubo scandiacus": (3, "day"),
    "Asio flammeus": (3, "dusk"),
    "Dendrocopos major": (1, None),
    "Corvus monedula": (1, None),
    "Corvus frugilegus": (1, None),
    "Corvus cornix": (1, None),
    "Corvus corone": (1, None),
    "Corvus corax": (3, None),
    "Parus major": (1, None),
    "Cyanistes caeruleus": (1, None),
    "Periparus ater": (1, None),
    "Sitta europaea": (1, None),
    "Erithacus rubecula": (1, None),
    "Turdus merula": (1, None),
    "Turdus torquatus": (3, None),
    "Phoenicurus ochruros": (1, None),
    "Cinclus cinclus": (2, None),
    "Passer domesticus": (1, None),
    "Passer montanus": (1, None),
    "Motacilla alba": (1, None),
    "Fringilla coelebs": (1, None),
    "Carduelis carduelis": (1, None),
    "Chloris chloris": (1, None),
    "Phalacrocorax carbo": (2, None),
}

NIGHT = "night"

# ---------------------------------------------------------------------------------------------------------------------
# Reading the data sets
# ---------------------------------------------------------------------------------------------------------------------


def die(message: str) -> None:
    print(f"build-species-info: {message}", file=sys.stderr)
    sys.exit(1)


def read_avonet(path: Path) -> tuple[dict[str, dict], dict[str, str]]:
    import openpyxl  # not needed by the rest of the project

    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = workbook["AVONET1_BirdLife"].iter_rows(values_only=True)
    head = next(rows)
    avonet = {}
    for row in rows:
        record = dict(zip(head, row))
        if record["Species1"]:
            avonet[record["Species1"]] = record
    rows = workbook["BirdLife-eBird crosswalk"].iter_rows(values_only=True)
    head = next(rows)
    to_birdlife: dict[str, str] = {}
    for row in rows:
        record = dict(zip(head, row))
        if record["Species1"] and record["Species2"]:
            to_birdlife.setdefault(record["Species2"], record["Species1"])
    return avonet, to_birdlife


def read_elton(path: Path) -> list[dict]:
    with path.open(encoding="latin-1", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def find(by_scientific: dict[str, dict], to_birdlife: dict[str, str], scientific: str) -> dict | None:
    if scientific in by_scientific:
        return by_scientific[scientific]
    birdlife = to_birdlife.get(scientific)
    return by_scientific.get(birdlife) if birdlife else None


# ---------------------------------------------------------------------------------------------------------------------
# The location model
# ---------------------------------------------------------------------------------------------------------------------


def weekly_scores(points: list[tuple[str, float, float]], indices: np.ndarray) -> np.ndarray:
    """Scores [point, week, bird] for the given BirdNET output positions."""
    from ai_edge_litert.interpreter import Interpreter

    interpreter = Interpreter(model_path=str(bs.GEO_MODEL), num_threads=2)
    inp = interpreter.get_input_details()[0]
    out = interpreter.get_output_details()[0]
    weeks = list(bs.WEEKS)
    interpreter.resize_tensor_input(inp["index"], [len(weeks), 3])
    interpreter.allocate_tensors()
    result = np.zeros((len(points), len(weeks), len(indices)), dtype=np.float32)
    for p, (_, lat, lon) in enumerate(points):
        interpreter.set_tensor(inp["index"], np.array([[lat, lon, week] for week in weeks], dtype=np.float32))
        interpreter.invoke()
        scores = np.round(interpreter.get_tensor(out["index"]).astype(np.float32), 5)
        result[p] = scores[:, indices]
    return result


def month_levels(weekly: np.ndarray) -> str:
    """weekly: 48 scores of one bird, averaged over its places. Twelve digits, January to December."""
    monthly = weekly.reshape(12, WEEKS_PER_MONTH).mean(axis=1)
    best = float(monthly.max())
    if best <= 0:
        return "0" * 12
    digits = []
    for value in monthly:
        share = float(value) / best
        digits.append(str(next((level for minimum, level in LEVEL_SHARES if share >= minimum), 0)))
    return "".join(digits)


# ---------------------------------------------------------------------------------------------------------------------


def build(args: argparse.Namespace) -> int:
    species = json.loads(SPECIES_FILE.read_text(encoding="utf-8"))
    if not bs.GEO_MODEL.exists():
        die(f"{bs.GEO_MODEL} not found. Run `node scripts/fetch-models.mjs` first.")

    # Tables must not name birds that are not in the Dex: a typo would silently do nothing.
    names = {s["scientificName"] for s in species}
    genera = {name.split(" ")[0] for name in names}
    problems = [f"SPECIES_RULES: {n}" for n in SPECIES_RULES if n not in names]
    problems += [f"PLACE_OVERRIDES: {n}" for n in PLACE_OVERRIDES if n not in names]
    problems += [f"GENUS_RULES: {g}" for g in GENUS_RULES if g not in genera]
    missing_points = [name for name, _, _ in bs.EUROPE_POINTS if name not in POINT_COUNTRY]
    if missing_points:
        problems.append(f"POINT_COUNTRY lacks {missing_points}")
    problems += [f"family {s['family']} has no FAMILY_RULES ({s['scientificName']})" for s in species if s["family"] not in FAMILY_RULES]
    if problems:
        die("\n  ".join(["fix the tables first:", *problems]))

    avonet, to_birdlife = read_avonet(Path(args.avonet))
    elton_rows = read_elton(Path(args.elton))
    elton_by_sci = {r["Scientific"]: r for r in elton_rows}
    elton_by_english = {r["English"].lower(): r for r in elton_rows}

    indices = np.array([s["birdnetIndex"] for s in species])
    print("running the location model at", len(bs.EUROPE_POINTS) + len(bs.CZECHIA_POINTS), "points ...", file=sys.stderr)
    europe = weekly_scores(bs.EUROPE_POINTS, indices)  # [point, week, bird]
    czechia = weekly_scores(bs.CZECHIA_POINTS, indices)
    europe_peak = europe.max(axis=1)  # [point, bird]
    country_of_point = np.array([POINT_COUNTRY[name] for name, _, _ in bs.EUROPE_POINTS])
    countries = sorted(set(country_of_point))

    unmatched: list[str] = []
    review: list[str] = []
    stats: dict[str, Counter] = defaultdict(Counter)
    for i, s in enumerate(species):
        sci = s["scientificName"]
        family = s["family"]
        shy, day, tip = FAMILY_RULES[family]
        genus = sci.split(" ")[0]
        if genus in GENUS_RULES:
            g_shy, g_day, g_tip = GENUS_RULES[genus]
            shy = g_shy if g_shy is not None else shy
            day = g_day if g_day is not None else day
            tip = g_tip if g_tip is not None else tip

        # Nocturnal birds are named by EltonTraits; a rule written for the species itself still wins.
        elton = find(elton_by_sci, to_birdlife, sci) or elton_by_english.get(s["englishName"].lower())
        if elton is None:
            unmatched.append(f"EltonTraits: {sci}")
        elif elton["Nocturnal"] == "1":
            day = NIGHT
        if sci in SPECIES_RULES:
            r_shy, r_day = SPECIES_RULES[sci]
            shy = r_shy if r_shy is not None else shy
            day = r_day if r_day is not None else day

        record = find(avonet, to_birdlife, sci)
        if record is None:
            die(f"AVONET has no row for {sci}")
        place = PLACE_OVERRIDES.get(sci) or PLACE_OF_HABITAT[record["Habitat"]]
        if s["inCzechia"] and sci not in PLACE_OVERRIDES:
            place = INLAND_PLACES.get(place, place)

        # Months: where the bird is in Czechia, or else over the part of Europe where the model puts it.
        if s["inCzechia"]:
            weekly = czechia[:, :, i].mean(axis=0)
        else:
            peak = europe_peak[:, i]
            chosen = np.flatnonzero(peak >= PRESENT_SCORE)
            if len(chosen) == 0:
                chosen = np.argsort(peak)[::-1][:RANGE_MIN_POINTS]
            weekly = europe[chosen, :, i].mean(axis=0)
        months = month_levels(weekly)
        if months == "0" * 12:
            die(f"{sci} has no month at all")

        info: dict = {"months": months, "place": place, "shy": shy, "day": day, "tip": tip}
        if not s["inCzechia"]:
            peak = europe_peak[:, i]
            by_country = {c: float(peak[country_of_point == c].max()) for c in countries}
            ranked = sorted(by_country.items(), key=lambda kv: -kv[1])
            floor = max(PRESENT_SCORE, COUNTRY_SHARE * ranked[0][1])
            found = [c for c, v in ranked if v >= floor]
            if not found:
                found = [c for c, v in ranked if v >= COUNTRY_FLOOR][:3]
            if not found:
                die(f"{sci} is in no country")
            info["countries"] = found[:MAX_COUNTRIES]
        s["info"] = info

        stats["place"][place] += 1
        stats["shy"][str(shy)] += 1
        stats["day"][day] += 1
        stats["tip"][tip] += 1
        review.append(
            f"| {s['dexNumber']} | {s['czechName'] or '-'} | {sci} | {family} | {'CZ' if s['inCzechia'] else 'EU'} | {months} | {place} | {shy} | {day} | {tip} | "
            + (",".join(info["countries"]) if "countries" in info else "")
            + " |"
        )

    for name, counter in stats.items():
        print(name, dict(counter.most_common()), file=sys.stderr)
    if unmatched:
        print("not in EltonTraits (daytime assumed unless the family rule says otherwise):", ", ".join(unmatched), file=sys.stderr)

    if args.review:
        head = (
            "# Species info for review\n\n"
            "Made by `scripts/build-species-info.py`. `months` is January to December (0 not there, 1 a few, 2 likely, 3 most likely);\n"
            "`shy` 1 = lets you come close, 2 = wary, 3 = shy or far away; `day` and `shy` are estimates by family and need a reader who knows birds.\n\n"
            "| # | Czech name | Latin | Family | Where | months | place | shy | day | tip | countries |\n|---|---|---|---|---|---|---|---|---|---|---|\n"
        )
        Path(args.review).write_text(head + "\n".join(review) + "\n", encoding="utf-8")

    if not args.no_write:
        text = "[\n" + ",\n".join(json.dumps(s, ensure_ascii=False, separators=(",", ":")) for s in species) + "\n]\n"
        SPECIES_FILE.write_text(text, encoding="utf-8")
        print(f"wrote {SPECIES_FILE} ({len(species)} birds)", file=sys.stderr)
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Add the finding info to public/data/species.json")
    parser.add_argument("--avonet", required=True, help="AVONET Supplementary dataset 1.xlsx")
    parser.add_argument("--elton", required=True, help="EltonTraits BirdFuncDat.txt")
    parser.add_argument("--review", help="also write one line per bird to this markdown file")
    parser.add_argument("--no-write", action="store_true", help="do everything but write species.json")
    return build(parser.parse_args())


if __name__ == "__main__":
    sys.exit(main())
